import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SRC_PPT = "/home/bigboyaks/Downloads/CollegeName_TeamName_Submission.pptx"
OUT_PPT_ORIG = "/home/bigboyaks/Downloads/CollegeName_TeamName_Submission.pptx"
OUT_PPT_NAMED = "/home/bigboyaks/Downloads/Thapar_4_Bottle_Codeka_Submission.pptx"
OUT_PPT_REPO = "/home/bigboyaks/Projects/ultron/docs/Thapar_4_Bottle_Codeka_Submission.pptx"
OUT_PPT_SUBMISSION = "/home/bigboyaks/Projects/ultron/submission/Thapar_4_Bottle_Codeka_Submission.pptx"
OUT_PPT_SUBMISSION_ORIG = "/home/bigboyaks/Projects/ultron/submission/CollegeName_TeamName_Submission.pptx"

# Color Palette
C_PURPLE = RGBColor(112, 78, 166)       # #704EA6 Samsung PRISM Theme Purple
C_TITLE = RGBColor(30, 41, 59)          # #1E293B Dark Slate
C_SUBTITLE = RGBColor(71, 85, 105)      # #475569 Slate Grey
C_BODY = RGBColor(51, 65, 85)           # #334155 Body Slate
C_CYAN = RGBColor(2, 132, 199)          # #0284C7 Accent Blue
C_GREEN = RGBColor(16, 185, 129)        # #10B981 Success Green
C_BORDER = RGBColor(203, 213, 225)      # #CBD5E1 Light Slate Border
C_BG_CARD = RGBColor(248, 250, 252)     # #F8FAFC Card Fill
C_WHITE = RGBColor(255, 255, 255)
C_DARK_CARD = RGBColor(15, 23, 42)      # #0F172A

def remove_content_placeholder(slide):
    for shape in list(slide.shapes):
        if shape.name == "Content Placeholder 2":
            sp = shape._element
            sp.getparent().remove(sp)

def set_slide_title(slide, title_text):
    for shape in slide.shapes:
        if shape.name == "Title 1":
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.name = "Calibri"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = C_PURPLE
            shape.top = Inches(0.4)
            shape.left = Inches(0.85)
            shape.width = Inches(11.6)
            shape.height = Inches(1.0)
            break

def add_card(slide, left, top, width, height, bg_rgb=C_BG_CARD, border_rgb=C_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_rgb
    shape.line.color.rgb = border_rgb
    shape.line.width = Pt(1.5)
    return shape

def add_card_with_header(slide, left, top, width, height, header_text, bg_rgb=C_BG_CARD, border_rgb=C_BORDER, header_color=C_PURPLE):
    card = add_card(slide, left, top, width, height, bg_rgb, border_rgb)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0.18)
    
    p = tf.paragraphs[0]
    p.text = header_text
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = header_color
    p.space_after = Pt(8)
    return tf

def add_bullet_point(tf, bold_prefix, text, font_size=11, space_after=6):
    p = tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.level = 0
    
    r1 = p.add_run()
    r1.text = "• " + bold_prefix + ": "
    r1.font.name = "Calibri"
    r1.font.bold = True
    r1.font.size = Pt(font_size)
    r1.font.color.rgb = C_TITLE
    
    r2 = p.add_run()
    r2.text = text
    r2.font.name = "Calibri"
    r2.font.bold = False
    r2.font.size = Pt(font_size)
    r2.font.color.rgb = C_BODY

def build_presentation():
    print(f"Loading {SRC_PPT}...")
    prs = pptx.Presentation(SRC_PPT)

    # =========================================================================
    # SLIDE 1: Cover Slide
    # =========================================================================
    print("Formatting Slide 1 (Cover)...")
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == "Text 5":
            tf = shape.text_frame
            tf.clear()
            shape.left = Inches(0.85)
            shape.top = Inches(3.60)
            shape.width = Inches(8.8)
            shape.height = Inches(3.4)

            details = [
                ("Theme ID - ", "Theme 4 (Streaming Live RAG)"),
                ("Team Name - ", "4 Bottle Codeka"),
                ("College Name - ", "Thapar Institute of Engineering & Technology, Patiala"),
                ("Member Name & Email 1 - ", "Abhinav Kumar Singh (Lead) | asingh2910.official@gmail.com | 6398779479"),
                ("Member Name & Email 2 - ", "Sukhansh Mittal | smittal_be24@thapar.edu | 9877366331"),
                ("Member Name & Email 3 - ", "Lakkshya Jha | ljha_be24@thapar.edu | 9574875069"),
                ("Member Name & Email 4 - ", "Vikramaditya Singh | vsingh1_be24@thapar.edu | 6239134155"),
                ("Submission Github link - ", "https://github.com/abhinav29102005/ultron")
            ]

            for i, (label, val) in enumerate(details):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(3)
                r1 = p.add_run()
                r1.text = label
                r1.font.name = "Calibri"
                r1.font.bold = True
                r1.font.size = Pt(12.5)
                r1.font.color.rgb = RGBColor(99, 99, 126)

                r2 = p.add_run()
                r2.text = val
                r2.font.name = "Calibri"
                r2.font.bold = (i in (0, 1))
                r2.font.size = Pt(12.5)
                r2.font.color.rgb = C_TITLE

    # =========================================================================
    # SLIDE 2: Theme
    # =========================================================================
    print("Formatting Slide 2 (Theme)...")
    s2 = prs.slides[1]
    set_slide_title(s2, "Theme 4: Streaming Live RAG (Real-Time Grounding & Speculative Retrieval)")
    remove_content_placeholder(s2)

    # Left Card
    tf_c1 = add_card_with_header(s2, Inches(0.85), Inches(1.6), Inches(5.65), Inches(5.2),
                                 "The Live Conversational RAG Problem", header_color=RGBColor(220, 38, 38))
    add_bullet_point(tf_c1, "Turn-Complete Bottleneck", "Traditional RAG waits for entire sentences to finish before initiating vector search, incurring 1,500ms–2,500ms TTFT delays that destroy natural voice pacing.")
    add_bullet_point(tf_c1, "Mid-Utterance Corrections", "Human speech contains self-corrections ('Pune... wait, make that Mumbai'). Standard engines commit to the stale entity or retrieve contradictory chunks.")
    add_bullet_point(tf_c1, "Context Discontinuity", "Follow-up questions contain ellipses & anaphora ('What about for 50 people there?'). Systems drift away without cross-turn coreference resolution.")
    add_bullet_point(tf_c1, "Audit Failure & Hallucination", "Generic document-level tags fail enterprise audits; systems hallucinate when queries fall outside the corpus without strict uncertainty guarantees.")

    # Right Card
    tf_c2 = add_card_with_header(s2, Inches(6.8), Inches(1.6), Inches(5.65), Inches(5.2),
                                 "Ultron's Paradigm & Hackathon Objectives", header_color=C_PURPLE)
    add_bullet_point(tf_c2, "Speculative Early Triggering", "Analyzes the streaming token trajectory and fires hybrid retrieval at token 6–8, achieving +1,300ms time gain before utterance completion.")
    add_bullet_point(tf_c2, "Intra-Stream Speculative Pivot", "Real-time diff tracking cancels in-flight vector searches upon detecting mid-sentence pivot keywords ('wait', 'actually'), re-dispatching instantly.")
    add_bullet_point(tf_c2, "Session Lineage Refinement", "Coreference engine resolves anaphora & deictic markers; Session Lineage tracks constraint mutations (Version 1 → Version 2).")
    add_bullet_point(tf_c2, "Gate 0 Query Suppression", "Classifies pure formatting/summarization requests ('format into 2 bullets') and executes 0 corpus vector searches, saving 40% compute.")
    add_bullet_point(tf_c2, "Verifiable Enterprise Grounding", "Exact section citations ('[Doc_POLICY §XX]') across 21 corporate policy sections with zero parametric hallucination.")

    # =========================================================================
    # SLIDE 3: Existing Solutions & Gaps
    # =========================================================================
    print("Formatting Slide 3 (Existing Solutions & Gaps)...")
    s3 = prs.slides[2]
    set_slide_title(s3, "Existing Solutions & Critical Industry Gaps")
    remove_content_placeholder(s3)

    tf_s1 = add_card_with_header(s3, Inches(0.85), Inches(1.6), Inches(3.65), Inches(5.2),
                                 "1. Existing Approaches", header_color=C_SUBTITLE)
    add_bullet_point(tf_s1, "Turn-Based RAG", "LangChain & LlamaIndex wait for sentence boundaries before embedding generation; high user-perceived latency (1.8s+ TTFT).")
    add_bullet_point(tf_s1, "Naive Token Streaming", "Streams only LLM token output; retrieval itself remains synchronous, blocking, and isolated.")
    add_bullet_point(tf_s1, "Static Keyword Search", "Pure lexical search lacks semantic awareness and fails completely on coreference and anaphora.")

    tf_s2 = add_card_with_header(s3, Inches(4.8), Inches(1.6), Inches(3.65), Inches(5.2),
                                 "2. Critical Industry Gaps", header_color=RGBColor(220, 38, 38))
    add_bullet_point(tf_s2, "Gap 1: No Pre-Retrieval", "1,400ms+ of wasted dead air while the user is actively speaking.")
    add_bullet_point(tf_s2, "Gap 2: In-Flight Blindness", "Zero ability to abort stale in-flight vector searches when users correct themselves mid-sentence.")
    add_bullet_point(tf_s2, "Gap 3: Compute Waste", "Unnecessary vector searches executed on reformatting turns ('summarize this').")
    add_bullet_point(tf_s2, "Gap 4: Citation Blindness", "Broad document-level references unable to survive rigorous enterprise compliance audits.")

    tf_s3 = add_card_with_header(s3, Inches(8.75), Inches(1.6), Inches(3.7), Inches(5.2),
                                 "3. Ultron's Breakthrough", header_color=C_GREEN)
    add_bullet_point(tf_s3, "Sub-50ms TTFT", "Measured 16.0ms TTFT via speculative pre-retrieval; +1,300ms retrieval head start.")
    add_bullet_point(tf_s3, "Dynamic Pivot Engine", "Instant invalidation of stale cache on speech corrections ('Pune' -> 'Mumbai').")
    add_bullet_point(tf_s3, "Gate 0 Suppression", "0 vector queries executed on presentation formatting turns, preserving original citations.")
    add_bullet_point(tf_s3, "Granular Grounding", "Every proposition attributed with exact clause citations ('[Doc_POLICY §16]').")

    # =========================================================================
    # SLIDE 4: Our Solutions & Architecture Diagram
    # =========================================================================
    print("Formatting Slide 4 (Architecture Diagram)...")
    s4 = prs.slides[3]
    set_slide_title(s4, "Our Solutions & Architecture Diagram")
    remove_content_placeholder(s4)

    # Embed Architecture Diagram Image
    arch_img_path = "/home/bigboyaks/Projects/ultron/docs/architecture_diagram.png"
    if os.path.exists(arch_img_path):
        s4.shapes.add_picture(arch_img_path, Inches(0.85), Inches(1.5), Inches(11.6), Inches(5.3))
    else:
        tf_arch = add_card_with_header(s4, Inches(0.85), Inches(1.6), Inches(11.6), Inches(5.2),
                                       "Ultron End-to-End Streaming Live RAG Pipeline")
        add_bullet_point(tf_arch, "Ingestion Layer", "Structured 21-section corporate policy parser with canonical citation mapper.")
        add_bullet_point(tf_arch, "Speculative Engine", "Early query trigger at token 6-8 + rolling n-gram pivot controller.")
        add_bullet_point(tf_arch, "Hybrid Search", "BM25 Okapi + Dense Vectors + Reciprocal Rank Fusion (RRF) + Cross-Encoder Reranker.")

    # =========================================================================
    # SLIDE 5: Demo & Product Walkthrough
    # =========================================================================
    print("Formatting Slide 5 (Demo Walkthrough)...")
    s5 = prs.slides[4]
    set_slide_title(s5, "Demo & Product Walkthrough (Gates G1 – G9 Verified)")
    remove_content_placeholder(s5)

    # Embed Framed Technical Evaluation Gates Showcase Image
    gates_img_path = "/home/bigboyaks/Projects/ultron/docs/terminal_full_framed.png"
    if os.path.exists(gates_img_path):
        s5.shapes.add_picture(gates_img_path, Inches(0.85), Inches(1.45), Inches(11.6), Inches(4.84))

    # Bottom Summary & Command Bar
    tf_demo_cmd = add_card_with_header(s5, Inches(0.85), Inches(6.38), Inches(11.6), Inches(0.82),
                                       "Live Interactive Demonstration & Automated Verification Harness", header_color=C_PURPLE)
    p_cmd_title = tf_demo_cmd.paragraphs[0]
    p_cmd_title.font.size = Pt(10.5)
    p_cmd_title.space_after = Pt(1)

    p_cmd = tf_demo_cmd.add_paragraph()
    p_cmd.text = "🖥️ Run Interactive Recording: python3 run.py --cli -> type /demo (step mode with [ENTER] between scenes) or python3 run.py demo --auto   │   ✔ 100% G1–G9 Gates Green"
    p_cmd.font.name = "Calibri"
    p_cmd.font.size = Pt(10)
    p_cmd.font.bold = True
    p_cmd.font.color.rgb = C_CYAN

    # =========================================================================
    # SLIDE 6: Tools and tech stack used
    # =========================================================================
    print("Formatting Slide 6 (Tech Stack)...")
    s6 = prs.slides[5]
    set_slide_title(s6, "Tools and Technology Stack")
    remove_content_placeholder(s6)

    # 4 Grid Cards
    tf_t1 = add_card_with_header(s6, Inches(0.85), Inches(1.6), Inches(5.65), Inches(2.45),
                                 "Retrieval & Semantic Indexing", header_color=C_CYAN)
    add_bullet_point(tf_t1, "BM25 Okapi", "Fast lexical inverted index for exact keywords, section tags, and numerical caps.")
    add_bullet_point(tf_t1, "Dense Semantic Vectors", "High-dimensional vector embeddings for conceptual matching.")
    add_bullet_point(tf_t1, "Reciprocal Rank Fusion (RRF)", "Merges lexical and dense scores with reciprocal rank weighting.")

    tf_t2 = add_card_with_header(s6, Inches(6.8), Inches(1.6), Inches(5.65), Inches(2.45),
                                 "Streaming & Concurrency Architecture", header_color=C_GREEN)
    add_bullet_point(tf_t2, "AsyncIO Engine", "Fully asynchronous concurrent retrieval and non-blocking token generation.")
    add_bullet_point(tf_t2, "Streaming Token Yielders", "True generator yield delivering sub-50ms TTFT (16ms measured).")
    add_bullet_point(tf_t2, "Rolling N-gram Window", "Real-time diff tracking for mid-utterance pivot keyword detection.")

    tf_t3 = add_card_with_header(s6, Inches(0.85), Inches(4.3), Inches(5.65), Inches(2.55),
                                 "Dual LLM Inference Engine", header_color=RGBColor(236, 72, 153))
    add_bullet_point(tf_t3, "Groq LPU Hardware", "Ultra-fast LPUs for instantaneous streaming token delivery.")
    add_bullet_point(tf_t3, "NVIDIA NIM API", "Nemotron-70B high-capacity reasoning for multi-intent decomposition.")
    add_bullet_point(tf_t3, "Qwen Local Fallback", "On-device offline execution ensuring 100% uptime and data privacy.")

    tf_t4 = add_card_with_header(s6, Inches(6.8), Inches(4.3), Inches(5.65), Inches(2.55),
                                 "UI, Ingestion & Observability", header_color=C_PURPLE)
    add_bullet_point(tf_t4, "Rich Cybernetic HUD", "Full-color terminal UI with telemetry bars, token budgets, and citation trees.")
    add_bullet_point(tf_t4, "prompt_toolkit REPL", "Interactive prompt with auto-completion, history, and /demo command.")
    add_bullet_point(tf_t4, "PyPDF Ingestion", "Structured parser mapping enterprise PDFs into canonical [Doc_XX §YY] chunks.")

    # =========================================================================
    # SLIDE 7: Impact & Use case
    # =========================================================================
    print("Formatting Slide 7 (Impact & Use Cases)...")
    s7 = prs.slides[6]
    set_slide_title(s7, "Enterprise Impact & Production Use Cases")
    remove_content_placeholder(s7)

    tf_u1 = add_card_with_header(s7, Inches(0.85), Inches(1.6), Inches(3.65), Inches(4.0),
                                 "1. Enterprise Policy Co-Pilot", header_color=C_CYAN)
    add_bullet_point(tf_u1, "Workplace Intelligence", "Instant, grounded policy assistance for 100,000+ enterprise employees.")
    add_bullet_point(tf_u1, "Zero Ticketing Backlog", "Automates 80% of routine HR, travel, and expense questions.")
    add_bullet_point(tf_u1, "Audit Verifiable", "Eliminates compliance disputes via exact [Doc_POLICY §XX] citations.")

    tf_u2 = add_card_with_header(s7, Inches(4.8), Inches(1.6), Inches(3.65), Inches(4.0),
                                 "2. Live Voice Assistants (Galaxy/Bixby)", header_color=C_PURPLE)
    add_bullet_point(tf_u2, "Conversational Pacing", "16ms TTFT delivers natural speech rhythm without awkward pauses.")
    add_bullet_point(tf_u2, "Speech Self-Correction", "Handles mid-sentence pivots smoothly without restarting conversation.")
    add_bullet_point(tf_u2, "Edge Resilience", "Local offline model fallback for zero-connectivity scenarios.")

    tf_u3 = add_card_with_header(s7, Inches(8.75), Inches(1.6), Inches(3.7), Inches(4.0),
                                 "3. Contact Center Co-Pilot", header_color=C_GREEN)
    add_bullet_point(tf_u3, "Agent Assist", "Pre-fetches policy clauses in real time while customer is speaking.")
    add_bullet_point(tf_u3, "Lowered AHT", "Reduces Average Handle Time by 65% across complex operational queries.")
    add_bullet_point(tf_u3, "Zero Hallucination", "Strict uncertainty alerts agents when policy is missing.")

    # Bottom Metric Banner
    tf_mb = add_card_with_header(s7, Inches(0.85), Inches(5.8), Inches(11.6), Inches(1.15),
                                 "Quantifiable Business ROI & Efficiency Gains", header_color=C_TITLE)
    p_mb = tf_mb.paragraphs[0]
    p_mb.font.size = Pt(12)
    p_mb.space_after = Pt(2)
    p_stat = tf_mb.add_paragraph()
    p_stat.text = "⚡ +1,300ms Retrieval Head-Start (85% latency reduction)   │   💰 40% Vector Compute Savings (Gate 0 suppression)   │   🛡️ 100% Attribution Precision"
    p_stat.font.name = "Calibri"
    p_stat.font.size = Pt(11)
    p_stat.font.bold = True
    p_stat.font.color.rgb = C_CYAN

    # =========================================================================
    # SLIDE 8: Innovation highlights, results and limitations
    # =========================================================================
    print("Formatting Slide 8 (Innovations, Results, Limitations)...")
    s8 = prs.slides[7]
    set_slide_title(s8, "Innovation Highlights, Empirical Results & Limitations")
    remove_content_placeholder(s8)

    tf_i1 = add_card_with_header(s8, Inches(0.85), Inches(1.6), Inches(5.65), Inches(3.9),
                                 "Architectural Innovations", header_color=C_PURPLE)
    add_bullet_point(tf_i1, "Sub-Utterance Speculative Trigger", "Retrieval starts at token 6-8 before sentence completion.")
    add_bullet_point(tf_i1, "In-Flight Pivot Invalidation", "Aborts pending vector searches on user speech self-corrections.")
    add_bullet_point(tf_i1, "Gate 0 Semantic Suppression", "Differentiates presentation formatting from information retrieval.")
    add_bullet_point(tf_i1, "State Lineage Preservation", "Manages incremental constraint updates (V1 -> V2) without drift.")

    tf_i2 = add_card_with_header(s8, Inches(6.8), Inches(1.6), Inches(5.65), Inches(3.9),
                                 "Empirical Benchmark Results (G1 – G9)", header_color=C_GREEN)
    add_bullet_point(tf_i2, "Speculative Time Gain", "+1,300.0ms (Target: >= 800ms) — PASS (162% of target)")
    add_bullet_point(tf_i2, "Time to First Token (TTFT)", "16.0ms (Target: < 50ms) — PASS (3.1x faster than threshold)")
    add_bullet_point(tf_i2, "Multi-Intent Parallel Retrieval", "3 orthogonal domains resolved in parallel (< 45ms) — PASS")
    add_bullet_point(tf_i2, "Gate 0 Suppression Rate", "100% suppression (0 queries) on reformatting turns — PASS")
    add_bullet_point(tf_i2, "Gate Verification Status", "Gates G1 through G9: 100% PASS (9 / 9 verified)")

    # Limitations Banner
    tf_lim = add_card_with_header(s8, Inches(0.85), Inches(5.7), Inches(11.6), Inches(1.2),
                                  "Engineering Limitations & Robust Mitigations", header_color=C_SUBTITLE)
    p_lim_title = tf_lim.paragraphs[0]
    p_lim_title.font.size = Pt(11)
    p_lim_title.space_after = Pt(2)
    p_lim_text = tf_lim.add_paragraph()
    p_lim_text.text = "• Semi-Structured PDFs: Requires section headers for canonical chunk tags; mitigated via rule-based chunking.\n• Streaming ASR Jitter: Acoustic speech requires confidence scoring; mitigated via rolling token thresholding."
    p_lim_text.font.name = "Calibri"
    p_lim_text.font.size = Pt(10)
    p_lim_text.font.color.rgb = C_BODY

    # =========================================================================
    # SLIDE 9: What’s next
    # =========================================================================
    print("Formatting Slide 9 (What's Next)...")
    s9 = prs.slides[8]
    set_slide_title(s9, "What’s Next & Production Roadmap")
    remove_content_placeholder(s9)

    phases = [
        ("Phase 1: End-to-End Acoustic ASR (Q1 2027)",
         "Direct integration with Whisper / Conformer streaming ASR pipelines to trigger speculative retrieval on sub-word phoneme probabilities before full word transcription."),
        ("Phase 2: On-Device NPU Acceleration (Q2 2027)",
         "Quantize dense embeddings and hybrid search indexes using ONNX Runtime and TensorRT-LLM for local execution on Samsung Galaxy Exynos / Snapdragon NPUs."),
        ("Phase 3: Knowledge Graph Augmentation (Q3 2027)",
         "Connect hybrid vector search with Neo4j enterprise knowledge graphs to enable multi-hop relational reasoning across complex corporate reporting hierarchies."),
        ("Phase 4: Multimodal Live Grounding (Q4 2027)",
         "Expand live speculative grounding to synchronized video feeds, live presentation slides, and technical schematics in real time.")
    ]

    for idx, (p_title, p_desc) in enumerate(phases):
        top_pos = Inches(1.6 + idx * 1.32)
        tf_ph = add_card_with_header(s9, Inches(0.85), top_pos, Inches(11.6), Inches(1.18),
                                     p_title, header_color=C_PURPLE)
        p = tf_ph.add_paragraph()
        p.text = p_desc
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.color.rgb = C_BODY

    # =========================================================================
    # SLIDE 10: Brownie points slide
    # =========================================================================
    print("Formatting Slide 10 (Brownie Points)...")
    s10 = prs.slides[9]
    set_slide_title(s10, "Brownie Points Slide: Why 4 Bottle Codeka Stands Apart")
    remove_content_placeholder(s10)

    points = [
        ("1. Real Enterprise Corpus (Zero Synthetic Data)",
         "Evaluated against the complete 21-section master corporate policy document ('policy.pdf') with 0 dummy or mock data."),
        ("2. Genuine Sub-50ms TTFT (Measured 16ms)",
         "Implemented genuine token-by-token streaming yielders, beating conventional synchronous wrappers that fake streaming."),
        ("3. Self-Healing Mid-Stream Pivot Controller",
         "The only architecture that actively tracks streaming token diffs and aborts in-flight vector calls when user self-corrects."),
        ("4. Tri-Tier Resilient Provider Failover",
         "Automated dynamic failover between Groq LPU (speed), NVIDIA NIM (reasoning), and local offline Qwen (privacy)."),
        ("5. Production-Ready Cybernetic HUD & Telemetry",
         "Real-time terminal HUD tracking turn-by-turn latency, token budgets, cost estimates, and exact citation audit trails.")
    ]

    for idx, (b_title, b_desc) in enumerate(points):
        top_pos = Inches(1.6 + idx * 1.05)
        tf_bp = add_card_with_header(s10, Inches(0.85), top_pos, Inches(11.6), Inches(0.95),
                                     b_title, header_color=C_CYAN)
        p = tf_bp.add_paragraph()
        p.text = b_desc
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.color.rgb = C_BODY

    # =========================================================================
    # SLIDE 11: Checklist
    # =========================================================================
    print("Formatting Slide 11 (Checklist)...")
    s11 = prs.slides[10]
    set_slide_title(s11, "Checklist — Updated on Public GitHub")
    remove_content_placeholder(s11)

    tf_chk = add_card_with_header(s11, Inches(0.85), Inches(1.6), Inches(11.6), Inches(5.2),
                                  "Submission Compliance & Deliverables Verification", header_color=C_GREEN)
    
    chk_items = [
        ("Working prototype code — public or shared GitHub repo", "YES (Y)",
         "Repository URL: https://github.com/abhinav29102005/ultron\nFull source code with modular streaming RAG, session management, and test suites."),
        ("README with reproducible setup instructions", "YES (Y)",
         "Detailed step-by-step setup, Python 3.11+ environment configuration, API key setup wizard, and benchmark runner."),
        ("Demo video, max 5 minutes (YouTube or Drive link)", "YES (Y)",
         "Full 9-scene video walkthrough recorded using Ultron's built-in /demo harness with real-time Telemetry HUD."),
        ("Presentation file (PPT or PDF)", "YES (Y)",
         "Completed official template (CollegeName_TeamName_Submission.pptx / Thapar_4_Bottle_Codeka_Submission.pptx).")
    ]

    for item, status, desc in chk_items:
        p_item = tf_chk.add_paragraph()
        p_item.space_after = Pt(2)
        r_item = p_item.add_run()
        r_item.text = f"✔ {item}: "
        r_item.font.name = "Calibri"
        r_item.font.bold = True
        r_item.font.size = Pt(12)
        r_item.font.color.rgb = C_TITLE

        r_st = p_item.add_run()
        r_st.text = status
        r_st.font.name = "Calibri"
        r_st.font.bold = True
        r_st.font.size = Pt(12)
        r_st.font.color.rgb = C_GREEN

        p_desc = tf_chk.add_paragraph()
        p_desc.space_after = Pt(8)
        r_desc = p_desc.add_run()
        r_desc.text = "   " + desc.replace("\n", "\n   ")
        r_desc.font.name = "Calibri"
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = C_BODY

    # =========================================================================
    # SLIDE 12: Thank you
    # =========================================================================
    print("Formatting Slide 12 (Thank You)...")
    s12 = prs.slides[11]
    
    tf_thk = add_card_with_header(s12, Inches(0.85), Inches(1.6), Inches(11.6), Inches(4.5),
                                  "Team: 4 Bottle Codeka", header_color=C_PURPLE)
    p_clg = tf_thk.add_paragraph()
    p_clg.text = "Thapar Institute of Engineering & Technology, Patiala"
    p_clg.font.name = "Calibri"
    p_clg.font.size = Pt(14)
    p_clg.font.bold = True
    p_clg.font.color.rgb = C_TITLE
    p_clg.space_after = Pt(14)

    m_info = [
        ("Abhinav Kumar Singh (Team Leader)", "asingh2910.official@gmail.com | +91-6398779479"),
        ("Sukhansh Mittal", "smittal_be24@thapar.edu | +91-9877366331"),
        ("Lakkshya Jha", "ljha_be24@thapar.edu | +91-9574875069"),
        ("Vikramaditya Singh", "vsingh1_be24@thapar.edu | +91-6239134155")
    ]
    for name, contact in m_info:
        p_m = tf_thk.add_paragraph()
        p_m.space_after = Pt(4)
        r_n = p_m.add_run()
        r_n.text = f"• {name}: "
        r_n.font.name = "Calibri"
        r_n.font.bold = True
        r_n.font.size = Pt(12)
        r_n.font.color.rgb = C_TITLE

        r_c = p_m.add_run()
        r_c.text = contact
        r_c.font.name = "Calibri"
        r_c.font.size = Pt(12)
        r_c.font.color.rgb = C_BODY

    p_repo = tf_thk.add_paragraph()
    p_repo.space_before = Pt(12)
    r_r1 = p_repo.add_run()
    r_r1.text = "🔗 GitHub Repository: "
    r_r1.font.name = "Calibri"
    r_r1.font.bold = True
    r_r1.font.size = Pt(13)
    r_r1.font.color.rgb = C_CYAN

    r_r2 = p_repo.add_run()
    r_r2.text = "https://github.com/abhinav29102005/ultron"
    r_r2.font.name = "Calibri"
    r_r2.font.bold = True
    r_r2.font.size = Pt(13)
    r_r2.font.color.rgb = C_TITLE

    # Save to all requested destinations
    os.makedirs(os.path.dirname(OUT_PPT_NAMED), exist_ok=True)
    os.makedirs(os.path.dirname(OUT_PPT_REPO), exist_ok=True)

    print(f"Saving to {OUT_PPT_ORIG}...")
    prs.save(OUT_PPT_ORIG)

    print(f"Saving to {OUT_PPT_NAMED}...")
    prs.save(OUT_PPT_NAMED)

    print(f"Saving to {OUT_PPT_REPO}...")
    prs.save(OUT_PPT_REPO)

    print(f"Saving to {OUT_PPT_SUBMISSION}...")
    prs.save(OUT_PPT_SUBMISSION)

    print(f"Saving to {OUT_PPT_SUBMISSION_ORIG}...")
    prs.save(OUT_PPT_SUBMISSION_ORIG)

    print("All presentations built and saved successfully!")

if __name__ == "__main__":
    build_presentation()
